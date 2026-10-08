import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import json

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(BASE_DIR, "agrash_schools_sheet.xlsx")

# Exact data list of all 58 schools
schools_data = [
    ["Delhi Public School DPS Kathua", "YES", "D-21", "", "", "YES", "DC-21", "YES", "SE-21", "F-501"],
    ["SMR International School Safidon", "YES", "D-22", "", "", "YES", "DC-22", "", "", "F-502"],
    ["ST Andrews World School Kundli", "", "", "YES", "S-23", "", "", "", "", "F-503"],
    ["Sant Nikka Singh Public School Model Town Karnal", "YES", "D-24", "YES", "S-24", "YES", "DC-24", "", "", "F-504"],
    ["JP Jain Sr Sec School Sonipat", "YES", "D-25", "", "", "YES", "DC-25", "", "", "F-505"],
    ["Himgiri Public School Samalkha", "YES", "D-26", "", "", "", "", "", "", "F-506"],
    ["Delhi Public School DPS Panipat City", "", "", "", "", "YES", "DC-27", "YES", "SE-27", "F-507"],
    ["Doon Public School Gohana", "YES", "D-28", "", "", "YES", "DC-28", "YES", "SE-28", "F-508"],
    ["Navyug Public High School", "", "", "", "", "", "", "YES", "SE-29", "F-509"],
    ["Arya Kanya Gurukul Sr Sec School Karnal", "YES", "D-30", "YES", "S-30", "YES", "DC-30", "", "", "F-510"],
    ["Pratap Public School Sector 6 Karnal", "", "", "", "", "YES", "DC-31", "YES", "SE-31", "F-511"],
    ["DAV Public School Thermal Colony Panipat", "YES", "D-32", "", "", "", "", "YES", "SE-32", "F-513"],
    ["DAV Police Public Sr Sec School Police Line Sonipat", "YES", "D-33", "", "", "", "", "YES", "SE-33", "F-514"],
    ["Guru Teg Bhadur Public School Karnal", "YES", "D-34", "YES", "S-34", "YES", "DC-34", "YES", "SE-34", "F-401"],
    ["Ujala Modern Sr Sec School Kaith Shahpur Israna", "YES", "D-35", "", "", "", "", "", "", "F-405"],
    ["Motilal Nehru Public School Urban Estate Jind", "YES", "D-36", "", "", "YES", "DC-36", "YES", "SE-36", "F-406"],
    ["DAV Police Public School Police Line Karnal", "YES", "D-37", "YES", "S-37", "YES", "DC-37", "YES", "SE-37", "F-407"],
    ["Sarv Vidya Public School Indri", "YES", "D-38", "", "", "YES", "DC-38", "", "", "F-408"],
    ["Halwasiya Vidya Vihar Sr Sec School Bhiwani", "YES", "D-39", "YES", "S-39", "YES", "DC-39", "YES", "SE-39", "F-409"],
    ["Shiv Shakti Sr Sec School Sonipat", "", "", "", "", "YES", "DC-40", "", "", "F-410"],
    ["Pratap Singh Memorial Sr Sec School Sonipat", "", "", "", "", "YES", "DC-41", "YES", "SE-41", "F-302"],
    ["Babu Ram ( BR ) Global School Gannaur", "", "", "YES", "S-42", "YES", "DC-42", "YES", "SE-42", "F-303"],
    ["MM Public School Assandh", "YES", "D-43", "YES", "S-43", "YES", "DC-43", "YES", "SE-43", "F-310"],
    ["Mothers Pride School Pataudi Gurugram", "YES", "D-44", "YES", "S-44", "YES", "DC-44", "YES", "SE-44", "F-311"],
    ["BR International Public School Kurukshetra", "YES", "D-45", "YES", "S-45", "YES", "DC-45", "YES", "SE-45", "F-202"],
    ["RPS International School Sector 50 Gurgaon", "", "", "", "", "", "", "YES", "SE-46", "F-203"],
    ["Holy Family Convent School Gohana", "", "", "", "", "YES", "DC-47", "", "", "F-204"],
    ["Mahavati Sports Sr Sec School Pawti Samalkha", "", "", "", "", "YES", "DC-48", "YES", "SE-48", "F-205"],
    ["Lakshay International School Naultha Panipat", "YES", "D-49", "", "", "YES", "DC-49", "YES", "SE-49", "F-206"],
    ["Rishikul Vidyapeeth Sonipat", "YES", "D-50", "YES", "S-50", "YES", "DC-50", "", "", "F-207"],
    ["Sant Nischal Singh Public School Yamunanagar", "YES", "D-51", "YES", "S-51", "YES", "DC-51", "YES", "SE-51", "F-209"],
    ["SD Sr Sec School Panipat", "YES", "D-52", "YES", "S-52", "YES", "DC-52", "YES", "SE-52", "F-210"],
    ["Geeta Vidya Mandir Public School Nimbari", "", "", "YES", "S-53", "", "", "YES", "SE-53", "D-201"],
    ["Dayanand Public School Gohana", "", "", "", "", "YES", "DC-54", "YES", "SE-54", "D-202"],
    ["Saraswati Public School Jagadhri Yamunanagar", "YES", "D-55", "YES", "S-55", "YES", "DC-55", "YES", "SE-55", "D-204"],
    ["DAV Centenary Public School Police Line Samalkha Panipat", "YES", "D-56", "YES", "S-56", "", "", "", "", "D-205"],
    ["JPS Academy Assandh", "YES", "D-57", "YES", "S-57", "YES", "DC-57", "YES", "SE-57", "D-207"],
    ["Delhi Public School Jind", "YES", "D-58", "YES", "S-58", "YES", "DC-58", "", "", "D-404"],
    ["Maharishi Dayanand Sr Sec School Ballah Karnal", "YES", "D-59", "", "", "", "", "", "", "D-405"],
    ["Yaduvanshi Niketan School Sector 33 Gurugram", "YES", "D-60", "", "", "YES", "DC-60", "YES", "SE-60", "E-303"],
    ["Sant Nikka Singh Public School Zarifa Farm Karnal", "YES", "D-61", "YES", "S-61", "YES", "DC-61", "", "", "E-304"],
    ["Geeta Vidya Mandir Public School NHBC Panipat", "YES", "D-62", "YES", "S-62", "YES", "DC-62", "", "", "E-305"],
    ["Shaheed Bhagat Singh Sr Sec School Karnal", "YES", "D-63", "", "", "YES", "DC-63", "", "", "E-307"],
    ["Swastik Bal Vikas Sr Sec School Panipat", "YES", "D-64", "YES", "S-64", "YES", "DC-64", "YES", "SE-64", "E-308"],
    ["Silver Bells Public School Shamli", "", "", "YES", "S-65", "YES", "DC-65", "", "", "E-310"],
    ["The Golden Era Public School", "YES", "D-66", "", "", "", "", "", "", ""],
    ["MASD Public School Panipat", "", "", "", "", "YES", "DC-67", "", "", ""],
    ["SD Modern Sr Sec School Panipat", "YES", "D-68", "", "", "YES", "DC-68", "", "", ""],
    ["DR MKK Arya Model Sr Sec School Panipat", "YES", "D-69", "YES", "S-69", "YES", "DC-69", "YES", "SE-69", ""],
    ["Global Public School Gohana", "", "", "", "", "YES", "DC-70", "", "", ""],
    ["Sunrise Public School Panipat", "YES", "D-71", "", "", "YES", "DC-71", "", "", ""],
    ["Gyandeep Sr Sec School Gannaur", "YES", "D-72", "", "", "YES", "DC-72", "YES", "SE-72", ""],
    ["OM Public School Gohana", "YES", "D-73", "YES", "S-73", "YES", "DC-73", "YES", "SE-73", ""],
    ["Montessori City School Kathua", "YES", "D-74", "", "", "", "", "", "", ""],
    ["Adarsh Vidya Mandir Sr Sec School Panipat", "YES", "D-75", "", "", "", "", "", "", ""],
    ["North Valley Public School", "YES", "D-76", "YES", "S-76", "YES", "DC-76", "", "", ""],
    ["RKSD Public School Kaithal", "YES", "D-77", "YES", "S-77", "", "", "", "", ""],
    ["DAV Police Public School Kurukshetra", "YES", "D-78", "YES", "S-78", "YES", "DC-78", "", "", ""]
]

# Styling Palette
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def build_workbook(scores_data=None):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # SHEET 1: Master Registration (All 58 Schools)
    # -------------------------------------------------------------
    ws_master = wb.active
    ws_master.title = "Master Registrations"
    ws_master.sheet_properties.tabColor = "00E5FF" # Cyan

    # Title
    ws_master.merge_cells("A1:J1")
    title_cell = ws_master["A1"]
    title_cell.value = "🔥 AGRASH GRAND FINALE - MASTER REGISTRATION & SLOTS"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="0B101C")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_master.row_dimensions[1].height = 32

    # Headers
    headers = ["School Name", "Group Dance", "Tags", "Group Song", "Tags", "Declamation", "Tags", "Science Exhibition", "Tags", "Green Room Allocation"]
    ws_master.append(headers)
    ws_master.row_dimensions[2].height = 24

    for col_idx in range(1, 11):
        cell = ws_master.cell(row=2, column=col_idx)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Summary Row
    summary = ["Total: 58 Schools", "43 Teams", "", "27 Teams", "", "44 Teams", "", "30 Teams", "", "Total: 144 Slots"]
    ws_master.append(summary)
    for col_idx in range(1, 11):
        cell = ws_master.cell(row=3, column=col_idx)
        cell.font = Font(name="Arial", size=9, bold=True, color="00E5FF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    for row_idx, r in enumerate(schools_data, start=4):
        ws_master.append(r)
        ws_master.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        row_bg = "F8FAFC" if is_even else "FFFFFF"

        for col_idx in range(1, 11):
            cell = ws_master.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Arial", size=9.5)
            cell.border = thin_border
            cell.fill = PatternFill("solid", fgColor=row_bg)
            
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Highlight Tags
            if col_idx in [3, 5, 7, 9] and cell.value:
                cell.font = Font(name="Arial", size=9.5, bold=True, color="FF2A4B")
            elif col_idx in [2, 4, 6, 8] and cell.value == "YES":
                cell.font = Font(name="Arial", size=9, bold=True, color="10B981")
            elif col_idx == 10 and cell.value:
                cell.font = Font(name="Arial", size=9.5, bold=True, color="2563EB")

    auto_fit_columns(ws_master)

    # -------------------------------------------------------------
    # HELPER TO BUILD CATEGORY SHEET
    # -------------------------------------------------------------
    def add_category_sheet(sheet_title, tab_color, tag_prefix, yes_col_idx, tag_col_idx, cat_name):
        ws_cat = wb.create_sheet(title=sheet_title)
        ws_cat.sheet_properties.tabColor = tab_color

        ws_cat.merge_cells("A1:E1")
        c1 = ws_cat["A1"]
        c1.value = f"🏆 AGRASH 2026 - {cat_name.upper()} ROSTER & TAGS"
        c1.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        c1.fill = PatternFill("solid", fgColor="0B101C")
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws_cat.row_dimensions[1].height = 30

        cat_headers = ["Slot #", "Team Tag", "School Name", "Green Room", "Status"]
        ws_cat.append(cat_headers)
        ws_cat.row_dimensions[2].height = 22

        for col_i in range(1, 6):
            cell = ws_cat.cell(row=2, column=col_i)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        slot_num = 1
        for s in schools_data:
            has_part = s[yes_col_idx] == "YES" or s[tag_col_idx]
            tag = s[tag_col_idx]
            school_name = s[0]
            room = s[9]

            if has_part and tag:
                r_idx = slot_num + 2
                ws_cat.append([slot_num, tag, school_name, room or "-", "Registered"])
                ws_cat.row_dimensions[r_idx].height = 20
                is_even = (slot_num % 2 == 0)
                row_bg = "F8FAFC" if is_even else "FFFFFF"

                for col_i in range(1, 6):
                    cell = ws_cat.cell(row=r_idx, column=col_i)
                    cell.font = Font(name="Arial", size=9.5)
                    cell.border = thin_border
                    cell.fill = PatternFill("solid", fgColor=row_bg)
                    
                    if col_i == 2:
                        cell.font = Font(name="Arial", size=10, bold=True, color="FF2A4B")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_i == 3:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                
                slot_num += 1

        auto_fit_columns(ws_cat)

    # -------------------------------------------------------------
    # SHEET 2: Group Dance
    # -------------------------------------------------------------
    add_category_sheet("Group Dance", "FF2A4B", "D-", 1, 2, "Group Dance")

    # -------------------------------------------------------------
    # SHEET 3: Group Song
    # -------------------------------------------------------------
    add_category_sheet("Group Song", "10B981", "S-", 3, 4, "Group Song")

    # -------------------------------------------------------------
    # SHEET 4: Declamation
    # -------------------------------------------------------------
    add_category_sheet("Declamation", "F59E0B", "DC-", 5, 6, "Declamation")

    # -------------------------------------------------------------
    # SHEET 5: Science Exhibition
    # -------------------------------------------------------------
    add_category_sheet("Science Exhibition", "8B5CF6", "SE-", 7, 8, "Science Exhibition")

    # -------------------------------------------------------------
    # SHEET 6: Live Judgement Marksheet
    # -------------------------------------------------------------
    ws_scores = wb.create_sheet(title="Live Marksheet")
    ws_scores.sheet_properties.tabColor = "F59E0B"

    ws_scores.merge_cells("A1:H1")
    s1 = ws_scores["A1"]
    s1.value = "📊 AGRASH LIVE JUDGEMENT MARKSHEET & SCORES"
    s1.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    s1.fill = PatternFill("solid", fgColor="0B101C")
    s1.alignment = Alignment(horizontal="center", vertical="center")
    ws_scores.row_dimensions[1].height = 30

    score_headers = ["Rank", "Team Tag", "Judge 1", "Judge 2", "Judge 3", "Judge 4", "Judge 5", "Overall Avg (/100)"]
    ws_scores.append(score_headers)
    ws_scores.row_dimensions[2].height = 22

    for col_i in range(1, 9):
        cell = ws_scores.cell(row=2, column=col_i)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Load from scores.json if available
    scores_file = os.path.join(BASE_DIR, "data", "scores.json")
    all_scores = []
    if os.path.exists(scores_file):
        try:
            with open(scores_file, "r", encoding="utf-8") as f:
                all_scores = json.load(f)
        except Exception:
            pass

    # Unique tags
    unique_tags = []
    for s in schools_data:
        for t_idx in [2, 4, 6, 8]:
            tag = s[t_idx]
            if tag and tag not in unique_tags:
                unique_tags.append(tag)

    # Rank calculations
    tag_rows = []
    for tag in unique_tags:
        t_scores = [sc for sc in all_scores if sc.get("tag_no") == tag]
        totals = [sc["total"] for sc in t_scores if "total" in sc]
        avg = round(sum(totals) / len(totals), 2) if totals else 0
        tag_rows.append({"tag": tag, "totals": totals, "avg": avg, "count": len(totals)})

    tag_rows.sort(key=lambda x: x["avg"], reverse=True)

    for idx, r in enumerate(tag_rows, start=1):
        row_num = idx + 2
        rank_str = f"#{idx}" if r["count"] > 0 else "-"
        avg_str = r["avg"] if r["count"] > 0 else "Pending"
        
        j_scores = []
        for j_i in range(1, 6):
            if j_i - 1 < len(r["totals"]):
                j_scores.append(r["totals"][j_i - 1])
            else:
                j_scores.append("-")

        ws_scores.append([rank_str, r["tag"], j_scores[0], j_scores[1], j_scores[2], j_scores[3], j_scores[4], avg_str])
        ws_scores.row_dimensions[row_num].height = 20
        is_even = (idx % 2 == 0)
        row_bg = "F8FAFC" if is_even else "FFFFFF"

        for col_i in range(1, 9):
            cell = ws_scores.cell(row=row_num, column=col_i)
            cell.font = Font(name="Arial", size=9.5)
            cell.border = thin_border
            cell.fill = PatternFill("solid", fgColor=row_bg)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_i == 2:
                cell.font = Font(name="Arial", size=10, bold=True, color="FF2A4B")
            elif col_i == 8 and r["count"] > 0:
                cell.font = Font(name="Arial", size=10, bold=True, color="00E5FF")
                cell.fill = PatternFill("solid", fgColor="0F172A")

    auto_fit_columns(ws_scores)

    wb.save(OUTPUT_FILE)
    print(f"✅ Successfully built multi-sheet Excel workbook at: {OUTPUT_FILE}")
    print(f"   Sheets included:")
    for s in wb.sheetnames:
        print(f"     • {s}")
    return OUTPUT_FILE

if __name__ == "__main__":
    build_workbook()
